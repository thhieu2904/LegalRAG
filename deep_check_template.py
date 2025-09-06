from docx import Document
import zipfile
import re

def deep_check_template():
    template_path = 'd:/Personal/LegalRAG_OCR/rag_service/data/storage/collections/quy_trinh_cap_ho_tich_cap_xa/documents/DOC_001/templates/Khai_sinh_template.docx'
    
    print('=== DEEP TEMPLATE CHECK ===')
    print(f'File: {template_path}')
    
    # Method 1: Using python-docx
    print('\n--- Method 1: python-docx ---')
    try:
        doc = Document(template_path)
        print(f'Paragraphs: {len(doc.paragraphs)}')
        print(f'Tables: {len(doc.tables)}')
        
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if text:
                print(f'Para {i}: "{text}"')
        
        # Check tables
        for i, table in enumerate(doc.tables):
            print(f'Table {i}: {len(table.rows)} rows, {len(table.columns)} cols')
            for row_idx, row in enumerate(table.rows):
                for col_idx, cell in enumerate(row.cells):
                    cell_text = cell.text.strip()
                    if cell_text:
                        print(f'  Cell[{row_idx},{col_idx}]: "{cell_text}"')
                        
    except Exception as e:
        print(f'python-docx error: {e}')
    
    # Method 2: Raw XML analysis
    print('\n--- Method 2: Raw XML ---')
    try:
        with zipfile.ZipFile(template_path, 'r') as docx:
            content = docx.read('word/document.xml').decode('utf-8')
            
            # Extract all text elements
            text_matches = re.findall(r'<w:t[^>]*>([^<]+)</w:t>', content)
            print(f'Found {len(text_matches)} text elements')
            
            # Look for CCCD-related patterns
            cccd_patterns = []
            for text in text_matches:
                if any(keyword in text.lower() for keyword in ['tên', 'sinh', 'địa chỉ', 'cccd', 'giới', 'scan_']):
                    cccd_patterns.append(text)
            
            print(f'CCCD-related patterns: {len(cccd_patterns)}')
            for i, pattern in enumerate(cccd_patterns):
                print(f'  {i+1}: "{pattern}"')
                
            # Look for placeholder patterns
            placeholder_matches = re.findall(r'\{\{[^}]+\}\}', content)
            print(f'Placeholder patterns: {len(placeholder_matches)}')
            for placeholder in placeholder_matches:
                print(f'  Found: {placeholder}')
                
    except Exception as e:
        print(f'XML analysis error: {e}')

if __name__ == "__main__":
    deep_check_template()
