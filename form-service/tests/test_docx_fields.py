"""Offline regression tests. All values are synthetic; no service is started."""
import asyncio
import hashlib
import importlib
import io
from pathlib import Path
import sys
import types
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.services.docx_fields import extract_docx_fields, fill_docx_fields


def packed(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def unpacked(content):
    return Document(io.BytesIO(content))


class DocxFieldsTests(unittest.TestCase):
    def test_plain_value_does_not_gain_a_dot(self):
        document = Document()
        document.add_paragraph('Số lượng: {{field_38}} bản')
        result = unpacked(fill_docx_fields(packed(document), {'field_38': '2'}))
        self.assertEqual(result.paragraphs[0].text, 'Số lượng: 2 bản')

    def test_split_runs_preserve_label_and_suffix_formatting(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run('Họ tên: ').bold = True
        paragraph.add_run('{{fi').italic = True
        paragraph.add_run('eld_2}}')
        paragraph.add_run(' — kết thúc').underline = True
        result = unpacked(fill_docx_fields(packed(document), {'field_2': 'TEN DEMO'})).paragraphs[0]
        self.assertEqual(result.text, 'Họ tên: TEN DEMO — kết thúc')
        self.assertTrue(result.runs[0].bold)
        self.assertIsNone(result.runs[1].bold)
        self.assertTrue(result.runs[1].italic)
        self.assertTrue(result.runs[-1].underline)

    def test_repeated_tokens_in_one_run(self):
        document = Document()
        document.add_paragraph('{{name}} và {{name}}; {{other}}')
        result = unpacked(fill_docx_fields(packed(document), {'name': 'DEMO', 'other': 'X'}))
        self.assertEqual(result.paragraphs[0].text, 'DEMO và DEMO; X')

    def test_values_are_literal_and_not_reprocessed(self):
        document = Document()
        document.add_paragraph('{{name}} {{other}}')
        value = r'C:\demo\1 {{other}}'
        result = unpacked(fill_docx_fields(packed(document), {'name': value, 'other': 'Y'}))
        self.assertEqual(result.paragraphs[0].text, value + ' Y')

    def test_nested_table(self):
        document = Document()
        nested = document.add_table(rows=1, cols=1).cell(0, 0).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = 'Số lượng {{field_38}}'
        content = packed(document)
        self.assertEqual(extract_docx_fields(content), ['field_38'])
        result = unpacked(fill_docx_fields(content, {'field_38': '3'}))
        self.assertEqual(result.tables[0].cell(0, 0).tables[0].cell(0, 0).text, 'Số lượng 3')

    def test_multiline_values_keep_word_breaks(self):
        document = Document()
        document.add_paragraph('{{left}} và {{right}}')
        result = unpacked(fill_docx_fields(packed(document), {'left': 'A\nB', 'right': 'C\tD'}))
        self.assertEqual(result.paragraphs[0].text, 'A\nB và C\tD')

    def test_existing_header_and_footer(self):
        document = Document()
        document.sections[0].header.paragraphs[0].text = '{{header}}'
        document.sections[0].footer.paragraphs[0].text = '{{footer}}'
        content = packed(document)
        self.assertEqual(set(extract_docx_fields(content)), {'header', 'footer'})
        result = unpacked(fill_docx_fields(content, {'header': 'H', 'footer': 'F'}))
        self.assertEqual(result.sections[0].header.paragraphs[0].text, 'H')
        self.assertEqual(result.sections[0].footer.paragraphs[0].text, 'F')

    def test_does_not_create_headers(self):
        document = Document()
        document.add_paragraph('{{field}}')
        content = packed(document)
        result = fill_docx_fields(content, {'field': 'DEMO'})
        with ZipFile(io.BytesIO(content)) as before, ZipFile(io.BytesIO(result)) as after:
            self.assertEqual(before.namelist(), after.namelist())

    def test_dotted_field_and_empty_field(self):
        document = Document()
        document.add_paragraph('{{name}}.... / {{empty}}')
        result = unpacked(fill_docx_fields(packed(document), {'name': 'A', 'empty': ''}))
        self.assertEqual(result.paragraphs[0].text, 'A' + '.' * 11 + ' / ' + '.' * 9)

    def test_hyperlink_run_is_filled(self):
        document = Document()
        paragraph = document.add_paragraph()
        link = OxmlElement('w:hyperlink')
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = '{{name}}'
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
        result = unpacked(fill_docx_fields(packed(document), {'name': 'DEMO'}))
        self.assertEqual(''.join(node.text or '' for node in result.paragraphs[0]._p.xpath('.//w:t')), 'DEMO')

    def test_drawing_in_edited_run_is_retained(self):
        document = Document()
        run = document.add_paragraph().add_run('{{name}}')
        run._r.append(OxmlElement('w:drawing'))
        result = unpacked(fill_docx_fields(packed(document), {'name': 'DEMO'}))
        self.assertEqual(len(result.element.xpath('.//w:drawing')), 1)
        self.assertEqual(result.paragraphs[0].text, 'DEMO')

    def test_actual_birth_template_all_37_fields(self):
        mounted = Path('/test-templates/birth.docx')
        path = mounted if mounted.is_file() else Path(__file__).resolve().parents[2] / 'thesis/refs/forms/Giấy đăng ký khai sinh.docx'
        original = path.read_bytes()
        self.assertEqual(hashlib.sha256(original).hexdigest(), '527768935ee88fa1e3b167c2a9300fc5b0a2ba7605cd79e75808d7e6db77b3c0')
        fields = extract_docx_fields(original)
        self.assertEqual(len(fields), 37)
        self.assertIn('form_38', fields)
        self.assertNotIn('form_37', fields)
        self.assertEqual({field for field in fields if field.startswith('scan_')},
                         {'scan_ho_ten', 'scan_ngay_sinh', 'scan_dia_chi', 'scan_cccd'})
        values = {field: f'DEMO-{field}' for field in fields}
        result = fill_docx_fields(original, values)
        self.assertEqual(extract_docx_fields(result), [])
        with ZipFile(io.BytesIO(original)) as before, ZipFile(io.BytesIO(result)) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(before.read(name), after.read(name), name)
        output = unpacked(result)
        requester = next(Paragraph(p, output) for p in output.element.xpath('.//w:p')
                         if 'DEMO-scan_ho_ten' in ''.join(node.text or '' for node in p.xpath('.//w:t')))
        value_run = next(run for run in requester.runs if 'DEMO-scan_ho_ten' in run.text)
        self.assertFalse(value_run.bold)

    def test_filler_refuses_changed_template(self):
        # Stub only unavailable network/config imports. Actual filling functions
        # are used; no storage, environment secrets or external data is read.
        http_stub = types.ModuleType('httpx')
        config_stub = types.ModuleType('config')
        config_stub.settings = types.SimpleNamespace(STORAGE_SERVICE_URL='unused', STORAGE_TIMEOUT=1)
        with patch.dict(sys.modules, {'httpx': http_stub, 'config': config_stub}):
            module = importlib.import_module('src.services.form_filler')
        filler = module.FormFiller()
        self.assertEqual(filler._prepare_context({'name': '  ', 'zero': 0}, ['name', 'zero']), {'name': '', 'zero': '0'})
        document = Document()
        document.add_paragraph('{{name}}')
        content = packed(document)

        async def download(_):
            return content

        filler._download_template = download
        result, error, _ = asyncio.run(filler.fill('fake.docx', {'name': 'DEMO'}, '0' * 64))
        self.assertIsNone(result)
        self.assertIn('thay đổi', error)
        result, error, validation = asyncio.run(filler.fill('fake.docx', {'name': 'DEMO'}, hashlib.sha256(content).hexdigest()))
        self.assertIsNone(error)
        self.assertEqual(validation['filled_fields'], 1)
        self.assertEqual(unpacked(result).paragraphs[0].text, 'DEMO')


class QRParserTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Test the actual parser only, not OpenCV/pyzbar image decoding.
        zbar_stub = types.ModuleType('pyzbar')
        zbar_stub.pyzbar = None
        with patch.dict(sys.modules, {'cv2': types.ModuleType('cv2'), 'pyzbar': zbar_stub}):
            cls.parser = importlib.import_module('src.services.cccd_scanner').QRCodeParser

    def test_valid_synthetic_payload(self):
        data = self.parser.parse_qr_data('000000000001||NGUOI DEMO|01011990|Nam|DIA CHI DEMO|01012021')
        self.assertTrue(self.parser.validate_cccd_data(data))
        self.assertEqual(data.field_ngay_sinh, '01/01/1990')
        self.assertIsNone(data.field_cmnd)

    def test_calendar_validation(self):
        self.assertIsNone(self.parser._format_date('31022022'))
        self.assertIsNone(self.parser._format_date('29022023'))
        self.assertEqual(self.parser._format_date('29022024'), '29/02/2024')

    def test_empty_reserved_columns_are_allowed(self):
        payload = '000000000001||NGUOI DEMO|01011990|Nam|DIA CHI DEMO|01012021'
        legacy = self.parser.parse_qr_data(payload)
        for suffix in ('|', '||||', '| | | | '):
            data = self.parser.parse_qr_data(payload + suffix)
            self.assertTrue(self.parser.validate_cccd_data(data))
            self.assertEqual(data, legacy)

    def test_nonempty_extensions_and_missing_core_columns_are_rejected(self):
        payload = '000000000001||NGUOI DEMO|01011990|Nam|DIA CHI DEMO|01012021'
        self.assertIsNone(self.parser.parse_qr_data(payload + '|UNKNOWN|||'))
        self.assertIsNone(self.parser.parse_qr_data(payload.rsplit('|', 1)[0]))

    def test_reject_malformed_id_and_date(self):
        self.assertIsNone(self.parser.parse_qr_data('not-an-id||DEMO|01011990|Nam|DEMO|01012021'))
        self.assertIsNone(self.parser.parse_qr_data('000000000001||DEMO|31021990|Nam|DEMO|01012021'))
        self.assertIsNone(self.parser.parse_qr_data('not a citizen identity payload'))


if __name__ == '__main__':
    unittest.main()
