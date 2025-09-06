#!/usr/bin/env python3
"""
Field Extraction Service - Extract tất cả {{placeholder}} từ Word files
Dùng để xây dựng Dynamic Form system
"""

import re
from pathlib import Path
from typing import Dict, List, Set, Optional, Any
import json
from docx import Document

class FieldExtractor:
    """
    Extract placeholder fields từ Word documents
    """
    
    def __init__(self):
        self.placeholder_pattern = re.compile(r'\{\{([^}]+)\}\}')
    
    def extract_fields_from_docx(self, docx_path: Path) -> Dict[str, Any]:
        """
        Extract tất cả {{placeholder}} từ .docx file
        
        Returns:
        {
            "total_fields": int,
            "scan_fields": [...],  # {{scan_*}} fields from CCCD
            "form_fields": [...],  # {{form_*}} fields cần user input
            "other_fields": [...], # Other {{placeholder}} patterns
            "field_metadata": {...} # Chi tiết về từng field
        }
        """
        
        if not docx_path.exists():
            raise FileNotFoundError(f"DOCX file not found: {docx_path}")
        
        # Read docx content
        doc = Document(str(docx_path))
        all_text = ""
        
        # Extract text from paragraphs
        for paragraph in doc.paragraphs:
            all_text += paragraph.text + "\n"
        
        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    all_text += cell.text + "\n"
        
        # Find all placeholder patterns
        matches = self.placeholder_pattern.findall(all_text)
        unique_fields = list(set(matches))  # Remove duplicates
        
        # Categorize fields
        scan_fields = [field for field in unique_fields if field.startswith('scan_')]
        form_fields = [field for field in unique_fields if field.startswith('form_')]
        other_fields = [field for field in unique_fields 
                       if not field.startswith('scan_') and not field.startswith('form_')]
        
        # Generate metadata for each field
        field_metadata = {}
        for field in unique_fields:
            field_metadata[field] = self._analyze_field(field, all_text)
        
        return {
            "total_fields": len(unique_fields),
            "scan_fields": sorted(scan_fields),
            "form_fields": sorted(form_fields), 
            "other_fields": sorted(other_fields),
            "field_metadata": field_metadata,
            "raw_fields": sorted(unique_fields)
        }
    
    def _analyze_field(self, field_name: str, content: str) -> Dict[str, Any]:
        """
        Analyze field để guess input type và validation rules
        """
        
        field_lower = field_name.lower()
        field_type = "text"  # default
        validation_rules = {}
        
        # Auto-detect field types based on naming
        if any(keyword in field_lower for keyword in ['ngay', 'date', 'thang', 'nam']):
            field_type = "date"
        elif any(keyword in field_lower for keyword in ['email', 'mail']):
            field_type = "email"
        elif any(keyword in field_lower for keyword in ['phone', 'dienthoai', 'sdt']):
            field_type = "tel"
        elif any(keyword in field_lower for keyword in ['so', 'number', 'cccd', 'cmnd']):
            field_type = "number"
        elif any(keyword in field_lower for keyword in ['gioitinh', 'gender']):
            field_type = "select"
            validation_rules["options"] = ["Nam", "Nữ"]
        elif any(keyword in field_lower for keyword in ['diachi', 'address']):
            field_type = "textarea"
        
        # Count occurrences in document
        pattern_count = content.count(f"{{{{{field_name}}}}}")
        
        # Generate display label
        display_label = self._generate_field_label(field_name)
        
        return {
            "field_name": field_name,
            "field_type": field_type,
            "display_label": display_label,
            "pattern_count": pattern_count,
            "validation_rules": validation_rules,
            "category": self._get_field_category(field_name)
        }
    
    def _generate_field_label(self, field_name: str) -> str:
        """
        Generate human-readable label cho field
        """
        
        # Remove prefixes
        label = field_name.replace('scan_', '').replace('form_', '')
        
        # Common translations
        translations = {
            'ho_ten': 'Họ và tên',
            'ngay_sinh': 'Ngày sinh', 
            'dia_chi': 'Địa chỉ',
            'gioi_tinh': 'Giới tính',
            'cccd': 'Số CCCD',
            'cmnd': 'Số CMND',
            'quanhe': 'Quan hệ',
            'tenkhaisinh': 'Tên khai sinh',
            'dt_khasinh': 'Ngày tháng khai sinh',
            'dantoc': 'Dân tộc',
            'quoctich': 'Quốc tịch',
            'noisinh': 'Nơi sinh',
            'quequan': 'Quê quán'
        }
        
        return translations.get(label, label.replace('_', ' ').title())
    
    def _get_field_category(self, field_name: str) -> str:
        """
        Categorize field for UI grouping
        """
        if field_name.startswith('scan_'):
            return 'cccd_data'
        elif field_name.startswith('form_'):
            return 'user_input'
        else:
            return 'other'


def test_field_extraction():
    """
    Test function để extract fields từ sample docx
    """
    
    # Test với file Khai_sinh.docx
    docx_path = Path("rag_service/data/storage/collections/quy_trinh_cap_ho_tich_cap_xa/documents/DOC_001/forms/Khai_sinh.docx")
    
    if not docx_path.exists():
        print(f"❌ Test file not found: {docx_path}")
        return
    
    extractor = FieldExtractor()
    
    try:
        result = extractor.extract_fields_from_docx(docx_path)
        
        print("🔍 FIELD EXTRACTION RESULTS")
        print("=" * 50)
        print(f"📊 Total fields found: {result['total_fields']}")
        print(f"🎯 CCCD fields (scan_*): {len(result['scan_fields'])}")
        print(f"📝 User input fields (form_*): {len(result['form_fields'])}")
        print(f"❓ Other fields: {len(result['other_fields'])}")
        print()
        
        print("🎯 CCCD Fields (auto-filled from scan):")
        for field in result['scan_fields']:
            metadata = result['field_metadata'][field]
            print(f"  • {field} → {metadata['display_label']} ({metadata['field_type']})")
        print()
        
        print("📝 User Input Fields (cần điền thủ công):")
        for field in result['form_fields']:
            metadata = result['field_metadata'][field]
            print(f"  • {field} → {metadata['display_label']} ({metadata['field_type']})")
        print()
        
        if result['other_fields']:
            print("❓ Other Fields:")
            for field in result['other_fields']:
                metadata = result['field_metadata'][field]
                print(f"  • {field} → {metadata['display_label']} ({metadata['field_type']})")
        
        # Save result to JSON for API testing
        output_file = Path("test_field_extraction_result.json")
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
        
        print(f"\n💾 Results saved to: {output_file}")
        
    except Exception as e:
        print(f"❌ Error during extraction: {e}")


if __name__ == "__main__":
    test_field_extraction()
