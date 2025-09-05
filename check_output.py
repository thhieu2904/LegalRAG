#!/usr/bin/env python
"""
Script kiểm tra nội dung file output
"""
from docx import Document

def check_output():
    output_path = 'test_output_fixed.docx'
    
    try:
        doc = Document(output_path)
        print('=== OUTPUT FILE CONTENT ===')
        
        all_text = ''
        for para in doc.paragraphs:
            if para.text.strip():
                all_text += para.text + '\n'
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        all_text += cell.text + '\n'
        
        # Check if data was filled
        test_values = [
            'Nguyen Van A',
            '01/01/1990', 
            '123 Nguyen Trai, Quan 1, TP.HCM',
            '123456789012',
            'Nam'
        ]
        
        print('--- DATA FILL CHECK ---')
        found_any = False
        for value in test_values:
            if value in all_text:
                print(f'✅ Found: {value}')
                found_any = True
            else:
                print(f'❌ Missing: {value}')
        
        if found_any:
            print('\n✅ SUCCESS: Template filling worked!')
        else:
            print('\n❌ FAILED: No data found in output')
            
        # Show sample output
        print('\n--- SAMPLE OUTPUT (first 3 lines with data) ---')
        count = 0
        for line in all_text.split('\n'):
            if any(val in line for val in test_values) and count < 3:
                print(f'  {line}')
                count += 1
                        
    except Exception as e:
        print(f'Error: {e}')

if __name__ == '__main__':
    check_output()
