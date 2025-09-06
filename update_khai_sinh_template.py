from docx import Document
from pathlib import Path

def update_khai_sinh_template():
    """Update Khai_sinh_template.docx với CCCD placeholders"""
    
    # Paths
    template_path = Path("d:/Personal/LegalRAG_OCR/rag_service/data/storage/collections/quy_trinh_cap_ho_tich_cap_xa/documents/DOC_001/templates/Khai_sinh_template.docx")
    backup_path = template_path.with_suffix('.backup.docx')
    
    print("=== UPDATING KHAI SINH TEMPLATE ===")
    print(f"Template path: {template_path}")
    
    try:
        # Backup original
        if template_path.exists():
            import shutil
            shutil.copy2(template_path, backup_path)
            print(f"✅ Backup created: {backup_path}")
        
        # Read original template
        doc = Document(str(template_path))
        print(f"Original paragraphs: {len(doc.paragraphs)}")
        
        # Add CCCD information section
        doc.add_paragraph()
        doc.add_heading('THÔNG TIN TỪ CCCD (Tự động điền)', level=2)
        doc.add_paragraph()
        
        # Add CCCD placeholders
        doc.add_paragraph(f'Họ và tên: {{ scan_ho_ten }}')
        doc.add_paragraph(f'Ngày sinh: {{ scan_ngay_sinh }}')
        doc.add_paragraph(f'Giới tính: {{ scan_gioi_tinh }}')
        doc.add_paragraph(f'Địa chỉ: {{ scan_dia_chi }}')
        doc.add_paragraph(f'Số CCCD: {{ scan_cccd }}')
        doc.add_paragraph()
        
        # Add note
        doc.add_paragraph('* Thông tin trên được điền tự động từ việc quét CCCD')
        doc.add_paragraph()
        doc.add_paragraph('='*50)
        
        # Save updated template
        doc.save(str(template_path))
        print(f"✅ Template updated successfully")
        print(f"New paragraphs: {len(doc.paragraphs)}")
        
        # Verify placeholders
        doc_verify = Document(str(template_path))
        placeholders_found = []
        for para in doc_verify.paragraphs:
            if '{{' in para.text and '}}' in para.text:
                placeholders_found.append(para.text.strip())
        
        print(f"\n=== VERIFICATION ===")
        print(f"Placeholders found: {len(placeholders_found)}")
        for placeholder in placeholders_found:
            print(f"  ✅ {placeholder}")
            
        return True
        
    except Exception as e:
        print(f"❌ Error updating template: {e}")
        # Restore backup if exists
        if backup_path.exists():
            import shutil
            shutil.copy2(backup_path, template_path)
            print(f"↩️ Restored from backup")
        return False

if __name__ == "__main__":
    success = update_khai_sinh_template()
    if success:
        print("\n🎉 Template update completed successfully!")
    else:
        print("\n💥 Template update failed!")
