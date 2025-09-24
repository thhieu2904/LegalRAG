"""
Template Filling Service for CCCD data
Unified service combining template download and smart CCCD auto-filling
Integrates FieldExtractor and AutoFillService logic
"""

import io
import logging
import aiohttp
import tempfile
from typing import Optional
import re
from pathlib import Path
from typing import Dict, Any, Optional, List
from fastapi import HTTPException
from docx import Document

logger = logging.getLogger(__name__)

class TemplateFillingService:
    def __init__(self, rag_service_url: Optional[str] = None):
        from app.core.config import settings
        if rag_service_url is None:
            rag_service_url = settings.RAG_SERVICE_URL
        self.rag_service_url = rag_service_url
        self.placeholder_pattern = re.compile(r'\{\{([^}]+)\}\}')
        logger.info("TemplateFillingService initialized with direct mapping logic")
        
    async def download_template(self, collection_id: str, doc_id: str, template_name: Optional[str] = None) -> bytes:
        """
        Download template file from rag_service
        
        Args:
            collection_id: Collection ID
            doc_id: Document ID
            template_name: Optional specific template name
            
        Returns:
            Template file content as bytes
        """
        try:
            # Use forms endpoint instead of templates for better structure alignment
            if template_name:
                # Use forms download endpoint that matches physical directory structure
                url = f"{self.rag_service_url}/api/forms/file/{collection_id}/{doc_id}/{template_name}/download"
            else:
                # Fallback to default template
                url = f"{self.rag_service_url}/api/templates/application_template.docx"
            
            logger.info(f"Downloading template from: {url}")
            
            async with aiohttp.ClientSession() as session:
                async with session.get(url) as response:
                    if response.status == 200:
                        content = await response.read()
                        logger.info(f"Downloaded template: {len(content)} bytes")
                        return content
                    else:
                        error_text = await response.text()
                        raise HTTPException(
                            status_code=response.status, 
                            detail=f"Failed to download template: {error_text}"
                        )
                        
        except aiohttp.ClientError as e:
            logger.error(f"Network error downloading template: {e}")
            raise HTTPException(
                status_code=503, 
                detail=f"Cannot connect to rag_service: {str(e)}"
            )
        except Exception as e:
            logger.error(f"Error downloading template: {e}")
            raise HTTPException(
                status_code=500, 
                detail=f"Template download failed: {str(e)}"
            )
    
    def extract_placeholders_from_docx(self, docx_content: bytes) -> Dict[str, Any]:
        """
        Extract all {{placeholder}} from DOCX content using integrated FieldExtractor logic
        
        Returns:
        {
            "scan_fields": [...],  # {{scan_*}} fields from CCCD
            "form_fields": [...],  # {{form_*}} fields need user input  
            "all_fields": [...]    # All unique placeholders
        }
        """
        try:
            # Create temporary document from bytes
            temp_file = io.BytesIO(docx_content)
            doc = Document(temp_file)
            
            # Extract all text content
            all_text = ""
            for paragraph in doc.paragraphs:
                all_text += paragraph.text + "\n"
            
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        all_text += cell.text + "\n"
            
            # Find all placeholders
            matches = self.placeholder_pattern.findall(all_text)
            unique_fields = list(set(matches))
            
            # Categorize by prefix
            scan_fields = [field for field in unique_fields if field.startswith('scan_')]
            form_fields = [field for field in unique_fields if field.startswith('form_')]
            
            logger.info(f"📋 Extracted placeholders: {len(scan_fields)} scan, {len(form_fields)} form")
            
            return {
                "scan_fields": scan_fields,
                "form_fields": form_fields, 
                "all_fields": unique_fields
            }
            
        except Exception as e:
            logger.error(f"Error extracting placeholders: {e}")
            return {"scan_fields": [], "form_fields": [], "all_fields": []}

    def prepare_smart_cccd_context(self, input_data: Dict[str, Any], template_placeholders: Optional[List[str]] = None) -> Dict[str, Any]:
        """
        Smart context preparation using direct {{placeholder}} mapping
        No hard-coded field mapping needed - template placeholders are self-descriptive
        
        Args:
            input_data: Combined data from frontend (scan_xxx + form_xxx fields)
            template_placeholders: List of placeholders from template ({{scan_xxx}}, {{form_xxx}})
            
        Returns:
            Direct-mapped context for template filling
        """
        context = {}
        
        # Direct mapping: input field name = template placeholder name
        if template_placeholders:
            for placeholder in template_placeholders:
                # Direct lookup - no conversion needed
                if placeholder in input_data:
                    context[placeholder] = str(input_data[placeholder])
                else:
                    context[placeholder] = ""  # Empty if not provided
                    
            logger.info(f"🎯 Direct context mapping: {len([k for k, v in context.items() if v])} fields filled")
        else:
            # Fallback: use all input data as-is
            for key, value in input_data.items():
                context[key] = str(value) if value is not None else ""
            
            logger.info(f"🎯 Fallback context mapping: {len(context)} fields total")
        
        logger.debug(f"Context details: {context}")
        return context
    
    def prepare_cccd_context(self, input_data: Dict[str, Any]) -> Dict[str, Any]:
        """
        Legacy method - redirects to smart context preparation
        """
        return self.prepare_smart_cccd_context(input_data)
    
    async def fill_template_with_cccd_data(
        self, 
        collection_id: str, 
        doc_id: str, 
        combined_data: Dict[str, Any],
        template_name: Optional[str] = None
    ) -> bytes:
        """
        Main method: Download template and fill with combined scan + form data
        
        Args:
            collection_id: Collection ID
            doc_id: Document ID  
            combined_data: Combined data (scan_xxx + form_xxx fields)
            template_name: Optional specific template name
            
        Returns:
            Filled .docx file as bytes
        """
        try:
            # Step 1: Download template from rag_service
            template_content = await self.download_template(collection_id, doc_id, template_name)
            
            # Step 2: Extract placeholders from template for direct mapping
            placeholder_info = self.extract_placeholders_from_docx(template_content)
            logger.info(f"📋 Template analysis: {len(placeholder_info['scan_fields'])} scan fields, {len(placeholder_info['form_fields'])} form fields")
            
            # Step 3: Prepare direct context mapping (no hard-coded conversion)
            context = self.prepare_smart_cccd_context(combined_data, placeholder_info['all_fields'])
            
            # Step 4: Fill template with direct-mapped data
            filled_content = await self._fill_docx_template(template_content, context)
            
            return filled_content
            
        except HTTPException:
            raise
        except Exception as e:
            logger.error(f"Error in fill_template_with_cccd_data: {e}")
            raise HTTPException(
                status_code=500, 
                detail=f"Template filling failed: {str(e)}"
            )
    
    async def _fill_docx_template(self, template_content: bytes, context: Dict[str, Any]) -> bytes:
        """
        Fill DOCX template with context data using python-docx-template
        
        Args:
            template_content: Template file content as bytes
            context: Data to fill in template
            
        Returns:
            Filled document as bytes
        """
        try:
            # Import here to avoid issues if library not installed
            try:
                from docxtpl import DocxTemplate
            except ImportError:
                raise HTTPException(
                    status_code=500,
                    detail="python-docx-template not installed. Please install it first."
                )
            
            # Create template from bytes
            template_file = io.BytesIO(template_content)
            doc = DocxTemplate(template_file)
            
            # Fill template with context
            doc.render(context)
            
            # Save to bytes
            output_buffer = io.BytesIO()
            doc.save(output_buffer)
            output_buffer.seek(0)
            
            filled_content = output_buffer.getvalue()
            logger.info(f"Template filled successfully: {len(filled_content)} bytes")
            
            return filled_content
            
        except Exception as e:
            logger.error(f"Error filling DOCX template: {e}")
            raise HTTPException(
                status_code=500, 
                detail=f"DOCX template filling failed: {str(e)}"
            )
