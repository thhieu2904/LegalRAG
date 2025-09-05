#!/usr/bin/env python
"""
Script kiểm tra nội dung template
"""
from docx import Document

def check_template():
    template_path = 'rag_service/data/storage/collections/quy_trinh_cap_ho_tich_cap_xa/documents/DOC_001/templates/Khai_sinh_template.docx'
    
    try:
        doc = Document(template_path)
        print('=== TEMPLATE CONTENT ===')
        
        print('\n--- PARAGRAPHS ---')
        for i, para in enumerate(doc.paragraphs):
            if para.text.strip():
                print(f'P{i}: {para.text}')
        
        print('\n--- TABLES ---')
        for table_idx, table in enumerate(doc.tables):
            print(f'\nTable {table_idx}:')
            for row_idx, row in enumerate(table.rows):
                for cell_idx, cell in enumerate(row.cells):
                    if cell.text.strip():
                        print(f'  [{row_idx},{cell_idx}]: {cell.text}')
        
        # Check for specific placeholders
        print('\n--- PLACEHOLDER SEARCH ---')
        placeholders = ['scan_ho_ten', 'scan_ngay_sinh', 'scan_dia_chi', 'scan_cccd', 'scan_gioi_tinh']
        all_text = '\n'.join([para.text for para in doc.paragraphs])
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    all_text += '\n' + cell.text
        
        for placeholder in placeholders:
            if placeholder in all_text:
                print(f'✅ Found: {placeholder}')
            else:
                print(f'❌ Missing: {placeholder}')
                
        # Check for {{ }} style placeholders
        import re
        template_placeholders = re.findall(r'\{\{([^}]+)\}\}', all_text)
        if template_placeholders:
            print(f'\n--- JINJA PLACEHOLDERS FOUND ---')
            for ph in template_placeholders:
                print(f'{{{{ {ph.strip()} }}}}')
        else:
            print('\n❌ No {{ }} style placeholders found!')
                        
    except Exception as e:
        print(f'Error: {e}')

if __name__ == '__main__':
    check_template()
