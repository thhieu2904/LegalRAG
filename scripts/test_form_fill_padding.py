"""
Test form filler with dot-padding functionality
Standalone test - doesn't require form-service config
"""

import sys
import os
import io
import re
from docx import Document
from docx.shared import Pt


class DotPaddingFormFiller:
    """
    Standalone Form Filler with dot-padding.
    Extracted core logic for testing.
    """
    
    # Minimum dots to keep after value
    MIN_PADDING_DOTS = 6
    
    def fill_docx(self, template_bytes: bytes, data: dict) -> bytes:
        """Fill template with dot-padding preservation."""
        doc = Document(io.BytesIO(template_bytes))
        
        # Process paragraphs
        for para in doc.paragraphs:
            self._process_paragraph(para, data)
        
        # Process tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        self._process_paragraph(para, data)
        
        # Save to bytes
        output = io.BytesIO()
        doc.save(output)
        return output.getvalue()
    
    def _process_paragraph(self, para, data: dict):
        """Process a paragraph, replacing placeholders with dot-padding."""
        full_text = para.text
        
        # Find all placeholders
        pattern = r'\{\{(\w+)\}\}(\.+)?'
        matches = list(re.finditer(pattern, full_text))
        
        if not matches:
            return
        
        # Process each placeholder
        for match in matches:
            placeholder = match.group(0)  # {{field_id}}... or {{field_id}}
            field_id = match.group(1)     # field_id
            dots = match.group(2) or ""   # trailing dots if any
            
            if field_id not in data:
                continue
            
            value = str(data[field_id])
            
            # Calculate padding
            if dots:
                # Has dots - calculate how many to keep
                original_len = len(f"{{{{{field_id}}}}}{dots}")
                value_len = len(value)
                remaining_dots = max(self.MIN_PADDING_DOTS, original_len - value_len)
                replacement = value + ("." * remaining_dots)
            else:
                # No dots - just replace
                replacement = value
            
            # Do replacement in runs
            self._replace_in_runs(para, placeholder, replacement)
    
    def _replace_in_runs(self, para, old_text: str, new_text: str):
        """Replace text across runs while preserving formatting."""
        # Build text position map
        runs_data = []
        pos = 0
        for run in para.runs:
            runs_data.append({
                'run': run,
                'start': pos,
                'end': pos + len(run.text),
                'text': run.text
            })
            pos += len(run.text)
        
        full_text = para.text
        start_idx = full_text.find(old_text)
        
        if start_idx == -1:
            return
        
        end_idx = start_idx + len(old_text)
        
        # Find affected runs
        replacement_done = False
        for i, rd in enumerate(runs_data):
            run = rd['run']
            run_start = rd['start']
            run_end = rd['end']
            
            # Check if this run intersects with our target
            if run_end <= start_idx:
                # Before target - keep as is
                continue
            elif run_start >= end_idx:
                # After target - keep as is
                continue
            else:
                # This run overlaps with target
                if not replacement_done:
                    # First overlapping run - insert replacement here
                    before = run.text[:max(0, start_idx - run_start)]
                    after_in_run = run.text[max(0, end_idx - run_start):]
                    run.text = before + new_text + after_in_run
                    replacement_done = True
                else:
                    # Subsequent overlapping runs - clear the overlapping part
                    if run_start >= start_idx and run_end <= end_idx:
                        # Entire run is within target - clear it
                        run.text = ""
                    elif run_start < end_idx:
                        # Partial overlap at start
                        run.text = run.text[end_idx - run_start:]

# Test directory
TEST_DIR = os.path.join(os.path.dirname(__file__), '..', 'docs', 'thamkhao', 'form_thamkhao')

def test_fill_form():
    """Test filling a template with dot-padding"""
    print("=" * 70)
    print("TEST: Fill form with dot-padding")
    print("=" * 70)
    
    # Use the template we just created
    template_path = os.path.join(TEST_DIR, 'forms_goc_template_new.docx')
    output_path = os.path.join(TEST_DIR, 'test_filled_output.docx')
    
    if not os.path.exists(template_path):
        print(f"Template not found: {template_path}")
        return
    
    # Sample form data
    form_data = {
        "ho_ten": "Nguyễn Văn A",
        "nam_sinh": "1990",
        "quan_he_voi_nguoi_uoc_khai_sin": "Cha đẻ",
        "ho_ten_2": "Nguyễn Văn Bé",
        "nam_sinh_2": "2024",
        "ghi_bang_chu": "Hai nghìn không trăm hai mươi tư",
        "gioi_tinh": "Nam",
        "dan_toc": "Kinh",
        "quoc_tich": "Việt Nam",
        "que_quan": "TP. Hồ Chí Minh",
        "ho_ten_3": "Trần Thị B",
        "ho_ten_4": "Nguyễn Văn C",
        "so_dinh_danh": "079123456789",
        "so": "01",
        "ngay_sinh": "15",
        "thang": "06",
        "nam_sinh_3": "2024",
        "lam_tai": "UBND Phường X, Quận Y",
        "ngay_sinh_2": "20",
        "thang_2": "06",
        "nam_sinh_4": "2024"
    }
    
    # Fill the form
    filler = DotPaddingFormFiller()
    
    # Read template
    with open(template_path, 'rb') as f:
        template_bytes = f.read()
    
    # Fill
    result_bytes = filler.fill_docx(template_bytes, form_data)
    
    # Save result
    with open(output_path, 'wb') as f:
        f.write(result_bytes)
    
    print(f"✓ Filled form saved: {output_path}")
    
    # Verify - read back and check content
    doc = Document(output_path)
    text_content = '\n'.join([p.text for p in doc.paragraphs])
    
    print("\n--- Verification (checking filled values) ---")
    for key, value in list(form_data.items())[:5]:  # Check first 5
        if value in text_content:
            print(f"  ✓ {key}: '{value}' found")
        else:
            print(f"  ✗ {key}: '{value}' NOT found")
    
    print(f"\nOutput file: {output_path}")
    print("Please open the file to verify dot-padding is preserved!")

def test_fill_simple_demo():
    """Create a simple demo template and test filling"""
    print("\n" + "=" * 70)
    print("TEST: Simple demo with dot-padding verification")
    print("=" * 70)
    
    # Create a simple demo template
    demo_template = os.path.join(TEST_DIR, 'demo_template.docx')
    demo_filled = os.path.join(TEST_DIR, 'demo_filled.docx')
    
    doc = Document()
    
    # Add test content with placeholders and dots
    doc.add_paragraph("PHIẾU ĐĂNG KÝ")
    doc.add_paragraph("")
    doc.add_paragraph("Họ và tên: {{ho_ten}}.........................................")
    doc.add_paragraph("Năm sinh: {{nam_sinh}}....................")
    doc.add_paragraph("Địa chỉ: {{dia_chi}}................................................................")
    doc.add_paragraph("Số CCCD: {{so_cccd}}...........................")
    doc.add_paragraph("")
    doc.add_paragraph("Ngày {{ngay}} tháng {{thang}} năm {{nam}}")
    
    doc.save(demo_template)
    print(f"✓ Created demo template: {demo_template}")
    
    # Fill it
    filler = DotPaddingFormFiller()
    
    with open(demo_template, 'rb') as f:
        template_bytes = f.read()
    
    form_data = {
        "ho_ten": "Nguyễn Văn An",
        "nam_sinh": "1985",
        "dia_chi": "123 Đường ABC, Q1, TP.HCM",
        "so_cccd": "079123456789",
        "ngay": "25",
        "thang": "06",
        "nam": "2024"
    }
    
    result_bytes = filler.fill_docx(template_bytes, form_data)
    
    with open(demo_filled, 'wb') as f:
        f.write(result_bytes)
    
    print(f"✓ Filled form saved: {demo_filled}")
    
    # Read and show content
    doc_filled = Document(demo_filled)
    print("\n--- Filled content ---")
    for p in doc_filled.paragraphs:
        if p.text.strip():
            print(f"  {p.text}")

if __name__ == "__main__":
    test_fill_simple_demo()
    test_fill_form()
    print("\n" + "=" * 70)
    print("ALL FILL TESTS COMPLETE")
    print("=" * 70)
