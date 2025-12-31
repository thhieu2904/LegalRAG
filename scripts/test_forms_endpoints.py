"""
Test /forms/detect and /forms/finalize endpoints
"""

import sys
import os
import base64

# Add form-service to path for imports
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'form-service'))

from fastapi.testclient import TestClient
from main import app

# Test directory
TEST_DIR = os.path.join(os.path.dirname(__file__), '..', 'docs', 'thamkhao', 'form_thamkhao')

def test_detect_endpoint():
    """Test POST /forms/detect"""
    print("=" * 70)
    print("TEST 1: /forms/detect endpoint")
    print("=" * 70)
    
    client = TestClient(app)
    
    # Use hotich.docx
    test_file = os.path.join(TEST_DIR, 'hotich.docx')
    
    if not os.path.exists(test_file):
        print(f"❌ Test file not found: {test_file}")
        return False
    
    with open(test_file, 'rb') as f:
        files = {'file': ('hotich.docx', f, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
        response = client.post('/forms/detect', files=files)
    
    print(f"Status: {response.status_code}")
    
    if response.status_code != 200:
        print(f"❌ Failed: {response.text}")
        return False
    
    data = response.json()
    
    print(f"Success: {data['success']}")
    print(f"Total positions: {data['total_positions']}")
    print(f"\nFirst 5 positions:")
    for pos in data['positions'][:5]:
        print(f"  [{pos['index']}] para={pos['paragraph_index']}, type={pos['pattern_type']}, text={pos['text'][:30]}...")
    
    if data['total_positions'] > 0:
        print("✅ Detect endpoint works!")
        return True
    else:
        print("⚠️  No positions detected")
        return False


def test_finalize_endpoint():
    """Test POST /forms/finalize"""
    print("\n" + "=" * 70)
    print("TEST 2: /forms/finalize endpoint")
    print("=" * 70)
    
    client = TestClient(app)
    
    # Use forms gốc.docx
    test_file = os.path.join(TEST_DIR, 'forms gốc.docx')
    
    if not os.path.exists(test_file):
        print(f"❌ Test file not found: {test_file}")
        return False
    
    # First detect
    with open(test_file, 'rb') as f:
        files = {'file': ('forms_goc.docx', f, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
        detect_response = client.post('/forms/detect', files=files)
    
    if detect_response.status_code != 200:
        print(f"❌ Detect failed: {detect_response.text}")
        return False
    
    detect_data = detect_response.json()
    total = detect_data['total_positions']
    print(f"Detected {total} positions")
    
    # Select first 5 positions
    selected = [0, 1, 2, 3, 4] if total >= 5 else list(range(total))
    selected_str = ','.join(map(str, selected))
    
    print(f"Selecting positions: {selected}")
    
    # Finalize with selected positions
    with open(test_file, 'rb') as f:
        files = {'file': ('forms_goc.docx', f, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
        data = {'selected_indices': selected_str}
        finalize_response = client.post('/forms/finalize', files=files, data=data)
    
    print(f"Status: {finalize_response.status_code}")
    
    if finalize_response.status_code != 200:
        print(f"❌ Failed: {finalize_response.text}")
        return False
    
    result = finalize_response.json()
    
    print(f"Success: {result['success']}")
    print(f"Placeholders created: {result['placeholders']}")
    print(f"Template size: {len(result['template_content'])} bytes (base64)")
    
    # Decode and save template
    if result['template_content']:
        template_bytes = base64.b64decode(result['template_content'])
        output_path = os.path.join(TEST_DIR, 'test_finalized_template.docx')
        
        with open(output_path, 'wb') as f:
            f.write(template_bytes)
        
        print(f"✅ Template saved: {output_path}")
        print(f"✅ Finalize endpoint works!")
        return True
    else:
        print("❌ No template content returned")
        return False


def test_integration():
    """Test full workflow"""
    print("\n" + "=" * 70)
    print("TEST 3: Integration - Detect → Finalize → Verify")
    print("=" * 70)
    
    from docx import Document
    import re
    
    # Check if template was created
    template_path = os.path.join(TEST_DIR, 'test_finalized_template.docx')
    
    if not os.path.exists(template_path):
        print("❌ Template file not found")
        return False
    
    # Open and verify placeholders
    doc = Document(template_path)
    
    all_text = ""
    for para in doc.paragraphs:
        all_text += para.text + "\n"
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_text += cell.text + "\n"
    
    # Find placeholders
    placeholders = re.findall(r'\{\{([^}]+)\}\}', all_text)
    
    print(f"Found {len(placeholders)} placeholders in template:")
    for ph in placeholders[:10]:  # Show first 10
        print(f"  - {{{{{ph}}}}}")
    
    if len(placeholders) > 0:
        print("✅ Integration test passed!")
        return True
    else:
        print("❌ No placeholders found in template")
        return False


if __name__ == "__main__":
    results = []
    
    results.append(test_detect_endpoint())
    results.append(test_finalize_endpoint())
    results.append(test_integration())
    
    print("\n" + "=" * 70)
    print("SUMMARY")
    print("=" * 70)
    print(f"Passed: {sum(results)}/{len(results)}")
    
    if all(results):
        print("✅ All tests passed!")
    else:
        print("⚠️  Some tests failed")
