"""
Test edge cases for dot-padding form filler
Verify các trường hợp thực tế
"""

import io
from docx import Document


class SimpleDotPaddingFiller:
    """Simplified filler - focus on core logic"""
    
    MIN_DOTS = 3  # Minimum dots to keep
    
    def fill(self, doc: Document, data: dict) -> Document:
        """Fill document with data, preserving dot structure"""
        
        # Process paragraphs
        for para in doc.paragraphs:
            self._fill_paragraph(para, data)
        
        # Process tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        self._fill_paragraph(para, data)
        
        return doc
    
    def _fill_paragraph(self, para, data: dict):
        """Replace {{placeholder}}... with value..."""
        import re
        
        text = para.text
        
        # Pattern: {{field_id}} followed by optional dots
        pattern = r'\{\{(\w+)\}\}(\.+)?'
        
        def replacer(match):
            field_id = match.group(1)
            dots = match.group(2) or ""
            
            if field_id not in data:
                return match.group(0)  # Keep original if no data
            
            value = str(data[field_id])
            
            if dots:
                # Calculate remaining dots
                total_space = len(match.group(0))  # {{field}}....
                remaining = total_space - len(value)
                padding_dots = max(self.MIN_DOTS, remaining)
                return value + ("." * padding_dots)
            else:
                return value
        
        new_text = re.sub(pattern, replacer, text)
        
        if new_text != text:
            # Simple approach: clear all runs and set new text
            # This preserves paragraph formatting but loses run-level formatting
            self._set_paragraph_text(para, new_text)
    
    def _set_paragraph_text(self, para, new_text: str):
        """Set paragraph text while trying to preserve some formatting"""
        if not para.runs:
            para.add_run(new_text)
            return
        
        # Keep first run's formatting, set new text
        first_run = para.runs[0]
        
        # Clear all runs
        for run in para.runs[1:]:
            run.text = ""
        
        first_run.text = new_text


def test_case(name: str, template_lines: list, data: dict, expected_contains: list):
    """Run a test case"""
    print(f"\n{'='*60}")
    print(f"TEST: {name}")
    print(f"{'='*60}")
    
    # Create template
    doc = Document()
    for line in template_lines:
        doc.add_paragraph(line)
    
    # Show template
    print("Template:")
    for p in doc.paragraphs:
        print(f"  │ {p.text}")
    
    # Fill
    filler = SimpleDotPaddingFiller()
    filled_doc = filler.fill(doc, data)
    
    # Show result
    print(f"\nData: {data}")
    print("\nResult:")
    for p in filled_doc.paragraphs:
        print(f"  │ {p.text}")
    
    # Verify
    all_text = " ".join([p.text for p in filled_doc.paragraphs])
    
    print("\nVerification:")
    all_pass = True
    for expected in expected_contains:
        if expected in all_text:
            print(f"  ✅ Contains: '{expected}'")
        else:
            print(f"  ❌ Missing: '{expected}'")
            all_pass = False
    
    return all_pass


def main():
    results = []
    
    # ============================================
    # CASE 1: Normal case - value shorter than dots
    # ============================================
    results.append(test_case(
        "Value shorter than dots",
        ["Họ tên: {{ho_ten}}......................................"],
        {"ho_ten": "Nguyễn Văn A"},
        ["Nguyễn Văn A..."]  # Should have dots after
    ))
    
    # ============================================
    # CASE 2: Value LONGER than dots
    # ============================================
    results.append(test_case(
        "Value LONGER than dots",
        ["Địa chỉ: {{dia_chi}}.........."],
        {"dia_chi": "123 Đường Nguyễn Huệ, Phường Bến Nghé, Quận 1, TP.HCM"},
        ["123 Đường Nguyễn Huệ"]  # Value should appear, minimum dots
    ))
    
    # ============================================
    # CASE 3: No dots after placeholder
    # ============================================
    results.append(test_case(
        "No dots (date format)",
        ["Ngày {{ngay}} tháng {{thang}} năm {{nam}}"],
        {"ngay": "25", "thang": "12", "nam": "2024"},
        ["Ngày 25 tháng 12 năm 2024"]
    ))
    
    # ============================================
    # CASE 4: Multiple placeholders same line
    # ============================================
    results.append(test_case(
        "Multiple placeholders on same line",
        ["Giới tính: {{gioi_tinh}}...... Dân tộc: {{dan_toc}}......"],
        {"gioi_tinh": "Nam", "dan_toc": "Kinh"},
        ["Nam...", "Kinh..."]
    ))
    
    # ============================================
    # CASE 5: Very short value
    # ============================================
    results.append(test_case(
        "Very short value (1 char)",
        ["Số lượng: {{so_luong}}...................."],
        {"so_luong": "5"},
        ["5..."]  # Should have minimum dots
    ))
    
    # ============================================
    # CASE 6: Empty value
    # ============================================
    results.append(test_case(
        "Empty value",
        ["Ghi chú: {{ghi_chu}}...................."],
        {"ghi_chu": ""},
        ["..."]  # Should keep dots
    ))
    
    # ============================================
    # CASE 7: Value with special characters
    # ============================================
    results.append(test_case(
        "Value with special chars",
        ["Email: {{email}}...................."],
        {"email": "test@gmail.com"},
        ["test@gmail.com"]
    ))
    
    # ============================================
    # CASE 8: Missing data (no replacement)
    # ============================================
    results.append(test_case(
        "Missing data key",
        ["Họ tên: {{ho_ten}}......................................"],
        {},  # No data provided
        ["{{ho_ten}}"]  # Should keep placeholder
    ))
    
    # ============================================
    # SUMMARY
    # ============================================
    print("\n" + "="*60)
    print("SUMMARY")
    print("="*60)
    passed = sum(results)
    total = len(results)
    print(f"Passed: {passed}/{total}")
    
    if passed == total:
        print("✅ All tests passed!")
    else:
        print("⚠️ Some tests failed - review needed")


if __name__ == "__main__":
    main()
