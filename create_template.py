from docx import Document
from pathlib import Path

# Create a simple Word template
doc = Document()
doc.add_heading('ĐƠN XIN VIỆC', 0)

# Add some paragraphs with placeholders - FIXED with double braces
doc.add_paragraph()
doc.add_paragraph(f'Họ, chữ đệm, tên người yêu cầu: {{{{ scan_ho_ten }}}}')
doc.add_paragraph(f'Ngày, tháng, năm sinh: {{{{ scan_ngay_sinh }}}}')
doc.add_paragraph(f'Giới tính: {{{{ scan_gioi_tinh }}}}')
doc.add_paragraph(f'Nơi cư trú: {{{{ scan_dia_chi }}}}')
doc.add_paragraph(f'Số căn cước công dân: {{{{ scan_cccd }}}}')
doc.add_paragraph()
doc.add_paragraph('Kính gửi: Ban Giám đốc Công ty')
doc.add_paragraph('Tôi tên như trên, xin được ứng tuyển vào vị trí nhân viên tại công ty.')
doc.add_paragraph('Tôi cam kết tuân thủ đầy đủ các quy định của công ty.')
doc.add_paragraph()
doc.add_paragraph('Xin cảm ơn!')
doc.add_paragraph()
doc.add_paragraph('Ngày ... tháng ... năm ...')
doc.add_paragraph('Người làm đơn')
doc.add_paragraph('(Ký tên và ghi rõ họ tên)')

# Save template
template_path = Path('d:/Personal/LegalRAG_OCR/rag_service/data/storage/templates/application_template.docx')
doc.save(template_path)
print(f'Template created: {template_path}')
print(f'File exists: {template_path.exists()}')
print(f'File size: {template_path.stat().st_size} bytes')
