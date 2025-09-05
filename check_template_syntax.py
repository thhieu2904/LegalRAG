from docx import Document

# Check original template
print("=== ORIGINAL TEMPLATE CHECK ===")
try:
    doc = Document('d:/Personal/LegalRAG_OCR/rag_service/data/storage/templates/application_template.docx')
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip():
            print(f"Line {i+1}: {para.text}")
            
    # Check for placeholder syntax
    print("\n=== PLACEHOLDER SYNTAX CHECK ===")
    placeholders_found = []
    for para in doc.paragraphs:
        text = para.text
        if '{' in text:
            placeholders_found.append(text.strip())
    
    for ph in placeholders_found:
        print(f"Found placeholder: {ph}")
        if '{{' in ph and '}}' in ph:
            print("  ✅ Correct double brace syntax")
        elif '{' in ph and '}' in ph:
            print("  ⚠️ Single brace - may need to be double")
            
except Exception as e:
    print(f"Error reading template: {e}")
